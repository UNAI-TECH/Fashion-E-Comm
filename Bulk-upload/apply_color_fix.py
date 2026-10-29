import openpyxl
import json
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

FIXED_XLSX = r"D:\Aanya Fashion Site\Fashion-E-Comm\Bulk-upload\output\Aanya_Fashions_Meesho_Bulk_Catalogue_FIXED.xlsx"
OUTPUT_V2_XLSX = r"D:\Aanya Fashion Site\Fashion-E-Comm\Bulk-upload\output\Aanya_Fashions_Meesho_Bulk_Catalogue_FIXED_v2.xlsx"
OUTPUT_COLORS_XLSX = r"D:\Aanya Fashion Site\Fashion-E-Comm\Bulk-upload\output\Aanya_Fashions_Meesho_Bulk_Catalogue_COLOR_CORRECTED.xlsx"

# The 51 valid colors from Meesho Validation Sheet Column Y:
# Aqua Blue, Baby Blue, Beige, Black, Blue, Bronze, Brown, Copper, Coral, Cream,
# Dark Pink, Gajari, Gold, Green, Grey, Grey Melange, Indigo, Khaki, Lavender,
# Lemon Yellow, Maroon, Mehendi, Metallic, Mint Green, Multicolor, Mustard,
# Navy Blue, Nude, Olive, Orange, Pastel Blue, Pastel Brown, Pastel Orange,
# Pastel Pink, Pastel Red, Pastel Yellow, Peach, Pink, Purple, Rama, Rani,
# Red, Rust, Sage, Silver, Sky Blue, Teal, Tricolor, Turquoise, White, Yellow

# Column AQ valid colors for Blouse Color:
# Aqua Blue, Beige, Black, Blue, Brown, Coral, Cream, Gold, Green, Grey, Grey Melange,
# Khaki, Lavendar, Maroon, Mehendi, Metallic, Multicolor, Mustard, Navy Blue, Nude,
# Olive, Orange, Peach, Pink, Purple, Red, Rust, Silver, Teal, White, Yellow

# Mapping of each SKU to (Saree Body Color [Col Y], Blouse Color [Col AQ], Reason)
COLOR_MAPPINGS = {
    'AF-SAR-TEALGREE-001': ('Teal', 'Teal', 'Body is Teal Green; Maroon Red is border/pallu'),
    'AF-SAR-ROYALPUR-002': ('Purple', 'Purple', 'Body is Royal Purple; Blue/Teal is border/pallu'),
    'AF-SAR-PLUMPURP-003': ('Purple', 'Purple', 'Body is Plum Purple; Teal is border/pallu'),
    'AF-SAR-MAGENTAT-004': ('Rani', 'Pink', 'Body is Magenta/Rani pink; Teal Blue is border/pallu'),
    'AF-SAR-GOLDENYE-005': ('Yellow', 'Yellow', 'Body is Golden Yellow; Red is border/pallu'),
    'AF-SAR-GREENRED-006': ('Green', 'Green', 'Body is Dark Green; Red is border/pallu'),
    'AF-SAR-RUSTORAN-007': ('Rust', 'Rust', 'Body is Rust Orange with gold checks'),
    'AF-SAR-RASPBERR-008': ('Pink', 'Pink', 'Body is Raspberry Pink with gold checks'),
    'AF-SAR-PLUMPURP-009': ('Purple', 'Purple', 'Body is Plum Purple with gold checks'),
    'AF-SAR-ROSEPINK-010': ('Pink', 'Pink', 'Body is Rose Pink with gold checks'),
    'AF-SAR-MAUVEGOL-011': ('Lavender', 'Lavendar', 'Body is Mauve (Lavender in Meesho dropdown)'),
    'AF-SAR-TEALBLUE-012': ('Teal', 'Teal', 'Body is Teal Blue with gold checks'),
    'AF-SAR-NAVYBLUE-013': ('Navy Blue', 'Navy Blue', 'Body is Navy Blue with gold checks'),
    'AF-SAR-EMERALDG-014': ('Green', 'Green', 'Body is Emerald Green with gold checks'),
    'AF-SAR-ROYALPLU-015': ('Purple', 'Purple', 'Body is Royal Plum Purple with gold checks'),
    'AF-SAR-PARROTGR-016': ('Green', 'Black', 'Body is Parrot Green; Black is contrast border & blouse'),
    'AF-SAR-MAGENTAP-017': ('Pink', 'Black', 'Body is Magenta Pink; Black is contrast border & blouse'),
    'AF-SAR-CORALPIN-018': ('Peach', 'Peach', 'Body is Coral/Soft Pink; Purple is contrast border'),
    'AF-SAR-IVORYGOL-019': ('Cream', 'Maroon', 'Body is Ivory Gold; Maroon is contrast border & blouse'),
    'AF-SAR-CHAMPAGN-020': ('Gold', 'Maroon', 'Body is Champagne Gold; Maroon is contrast border & blouse'),
    'AF-SAR-AQUABLUE-021': ('Aqua Blue', 'Black', 'Body is Aqua Blue; Black is contrast blouse'),
    'AF-SAR-ROYALBLU-022': ('Blue', 'Blue', 'Body is Royal Blue with antique gold brocade'),
    'AF-SAR-PEACOCKT-023': ('Teal', 'Teal', 'Body is Peacock Teal Blue with antique gold zari'),
    'AF-SAR-ROYALPLU-024': ('Purple', 'Purple', 'Body is Royal Plum Purple with antique gold zari'),
    'AF-SAR-NAVYBLUE-025': ('Navy Blue', 'Navy Blue', 'Body is Navy Blue with antique gold zari'),
    'AF-SAR-MAROONRE-026': ('Maroon', 'Maroon', 'Body is Maroon Red with gold zari'),
    'AF-SAR-MUSTARDY-027': ('Mustard', 'Mustard', 'Body is Mustard Yellow with antique gold zari'),
    'AF-SAR-FUCHSIAP-028': ('Pink', 'Pink', 'Body is Fuchsia Pink with gold zari'),
    'AF-SAR-EMERALDG-029': ('Green', 'Green', 'Body is Emerald Green with antique gold checks'),
    'AF-SAR-DARKEMER-030': ('Green', 'Green', 'Body is Dark Emerald Green with antique gold zari'),
    'AF-SAR-ROYALBLU-031': ('Blue', 'Blue', 'Body is Royal Blue with antique gold zari checks'),
    'AF-SAR-TEALBLUE-032': ('Teal', 'Teal', 'Body is Deep Teal Blue with antique gold checks'),
    'AF-SAR-DARKBLAC-033': ('Purple', 'Purple', 'Body is Dark Blackish Purple with peacock zari'),
    'AF-SAR-ROYALPUR-034': ('Purple', 'Purple', 'Body is Royal Purple with elephant zari border'),
    'AF-SAR-ROYALBLU-035': ('Blue', 'Blue', 'Body is Royal Blueberry Blue with elephant zari'),
    'AF-SAR-ROYALBLU-036': ('Blue', 'Blue', 'Body is Royal Blue with gold zari border'),
    'AF-SAR-DARKTEAL-037': ('Teal', 'Teal', 'Body is Dark Teal Green with elephant zari border'),
    'AF-SAR-EMERALDG-038': ('Green', 'Green', 'Body is Emerald Green with gold zari border'),
    'AF-SAR-KANCHICO-039': ('Mustard', 'Mustard', 'Body is Mustard Yellow; Antique gold is zari border'),
    'AF-SAR-KANCHICO-040': ('Grey', 'Grey', 'Body is Grey; Black & Gold is zari border'),
    'AF-SAR-KANCHICO-041': ('Pink', 'Pink', 'Body is Fuchsia Pink; Deep Maroon & Gold is zari border'),
    'AF-SAR-KANCHICO-042': ('Peach', 'Peach', 'Body is Peach / Soft Pink; Maroon is zari border')
}

wb = openpyxl.load_workbook(FIXED_XLSX)
ws = wb['Sarees-Fill this']

report_rows = []
changes_count = 0

for r in range(5, ws.max_row + 1):
    name = ws.cell(row=r, column=4).value
    if not name:
        break
    sku = ws.cell(row=r, column=39).value
    old_color = ws.cell(row=r, column=25).value
    old_blouse = ws.cell(row=r, column=43).value
    
    if sku in COLOR_MAPPINGS:
        new_color, new_blouse, reason = COLOR_MAPPINGS[sku]
        
        ws.cell(row=r, column=25).value = new_color
        ws.cell(row=r, column=43).value = new_blouse
        
        changed = (old_color != new_color) or (old_blouse != new_blouse)
        if changed:
            changes_count += 1
            
        report_rows.append({
            'row': r,
            'sku': sku,
            'name': name,
            'old_color': old_color,
            'new_color': new_color,
            'old_blouse': old_blouse,
            'new_blouse': new_blouse,
            'changed': changed,
            'reason': reason
        })

# Save to output files
wb.save(OUTPUT_V2_XLSX)
print(f"Successfully saved to: {OUTPUT_V2_XLSX}")

wb.save(OUTPUT_COLORS_XLSX)
print(f"Successfully saved to: {OUTPUT_COLORS_XLSX}")

# Try to overwrite FIXED_XLSX if unlocked
try:
    wb.save(FIXED_XLSX)
    print(f"Successfully overwrote: {FIXED_XLSX}")
except PermissionError:
    print(f"Note: {FIXED_XLSX} is currently locked by Excel (open in background).")
    print("Saved copy to FIXED_v2.xlsx and COLOR_CORRECTED.xlsx!")

print("\n" + "=" * 110)
print(f"COLOR FIX SUMMARY: {changes_count} products updated with accurate saree body colors")
print("=" * 110)
print(f"{'#':2s} | {'SKU':20s} | {'Old Color':12s} -> {'New Color':12s} | {'Blouse Col':12s} | {'Status':8s} | {'Reason'}")
print("-" * 110)
for i, item in enumerate(report_rows, 1):
    status = "UPDATED" if item['changed'] else "KEPT"
    print(f"{i:2d} | {item['sku']:20s} | {str(item['old_color']):12s} -> {item['new_color']:12s} | {item['new_blouse']:12s} | {status:8s} | {item['reason']}")
