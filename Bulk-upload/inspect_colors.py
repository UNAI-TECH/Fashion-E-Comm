import openpyxl
import json
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'D:\Aanya Fashion Site\Fashion-E-Comm\Bulk-upload\output\Aanya_Fashions_Meesho_Bulk_Catalogue_FIXED.xlsx', data_only=True)
ws = wb['Sarees-Fill this']

with open(r'D:\Aanya Fashion Site\Fashion-E-Comm\Bulk-upload\output\scraped_products.json', 'r', encoding='utf-8') as f:
    scraped = json.load(f)

img_base = r'D:\Aanya Fashion Site\Fashion-E-Comm\Bulk-upload\output\images'

products = []
for r in range(5, ws.max_row + 1):
    name = ws.cell(row=r, column=4).value
    if not name:
        break
    sku = ws.cell(row=r, column=39).value
    current_col = ws.cell(row=r, column=25).value
    current_blouse = ws.cell(row=r, column=43).value
    
    p_scraped = None
    for p in scraped:
        if p.get('name') == name or p.get('name', '').replace('Coral', 'Soft Pink').replace('coral', 'soft pink') == name:
            p_scraped = p
            break
        elif sku and sku in p.get('sku', ''):
            p_scraped = p
            break
    
    spec_col = p_scraped.get('specifications', {}).get('color') if p_scraped else ''
    orig_name = p_scraped.get('name') if p_scraped else name
    desc = p_scraped.get('description', '') if p_scraped else ''
    
    sku_dir = os.path.join(img_base, sku) if sku else ''
    front_img = None
    if sku_dir and os.path.isdir(sku_dir):
        for fn in os.listdir(sku_dir):
            if 'image-1-front' in fn:
                front_img = os.path.join(sku_dir, fn)
                break
    
    products.append({
        'row': r,
        'sku': sku,
        'name': name,
        'orig_name': orig_name,
        'spec_color': spec_col,
        'current_col': current_col,
        'current_blouse': current_blouse,
        'front_img': front_img,
        'desc': desc
    })

print(f"Total products: {len(products)}")
for p in products:
    print(f"Row {p['row']:2d} | SKU: {p['sku']} | Current Col: {str(p['current_col']):12s} | Spec: {str(p['spec_color']):35s} | Name: {p['name']}")
