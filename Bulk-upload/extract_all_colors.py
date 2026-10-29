import openpyxl
import json
import os
import sys
from PIL import Image
import numpy as np
from collections import Counter

sys.stdout.reconfigure(encoding='utf-8')

# Valid Meesho colors from Validation Sheet Column Y:
# Aqua Blue, Baby Blue, Beige, Black, Blue, Bronze, Brown, Copper, Coral, Cream,
# Dark Pink, Gajari, Gold, Green, Grey, Grey Melange, Indigo, Khaki, Lavender,
# Lemon Yellow, Maroon, Mehendi, Metallic, Mint Green, Multicolor, Mustard,
# Navy Blue, Nude, Olive, Orange, Pastel Blue, Pastel Brown, Pastel Orange,
# Pastel Pink, Pastel Red, Pastel Yellow, Peach, Pink, Purple, Rama, Rani,
# Red, Rust, Sage, Silver, Sky Blue, Teal, Tricolor, Turquoise, White, Yellow

wb = openpyxl.load_workbook(r'D:\Aanya Fashion Site\Fashion-E-Comm\Bulk-upload\output\Aanya_Fashions_Meesho_Bulk_Catalogue_FIXED.xlsx', data_only=True)
ws = wb['Sarees-Fill this']

with open(r'D:\Aanya Fashion Site\Fashion-E-Comm\Bulk-upload\output\scraped_products.json', 'r', encoding='utf-8') as f:
    scraped = json.load(f)

img_base = r'D:\Aanya Fashion Site\Fashion-E-Comm\Bulk-upload\output\images'

items = []

for r in range(5, ws.max_row + 1):
    name = ws.cell(row=r, column=4).value
    if not name:
        break
    sku = ws.cell(row=r, column=39).value
    curr_color = ws.cell(row=r, column=25).value
    curr_blouse = ws.cell(row=r, column=43).value
    
    # Match scraped
    p_scraped = None
    for p in scraped:
        if p.get('name') == name or p.get('name', '').replace('Coral', 'Soft Pink').replace('coral', 'soft pink') == name:
            p_scraped = p
            break
        elif sku and sku in p.get('sku', ''):
            p_scraped = p
            break
            
    spec_col = p_scraped.get('specifications', {}).get('color', '') if p_scraped else ''
    orig_name = p_scraped.get('name') if p_scraped else name
    
    # Image analysis
    sku_dir = os.path.join(img_base, sku) if sku else ''
    front_img = None
    if sku_dir and os.path.isdir(sku_dir):
        for fn in sorted(os.listdir(sku_dir)):
            if 'image-1-front' in fn:
                front_img = os.path.join(sku_dir, fn)
                break
                
    top_colors = []
    if front_img and os.path.exists(front_img):
        im = Image.open(front_img).convert('RGB')
        arr = np.array(im)
        # remove white/near-white bg
        is_bg = (arr[:, :, 0] > 235) & (arr[:, :, 1] > 235) & (arr[:, :, 2] > 235)
        non_bg = arr[~is_bg]
        if len(non_bg) > 0:
            sample = non_bg[::40]
            quant = [(int(p[0]//20)*20, int(p[1]//20)*20, int(p[2]//20)*20) for p in sample]
            top_colors = Counter(quant).most_common(5)
            
    items.append({
        'row': r,
        'sku': sku,
        'name': name,
        'orig_name': orig_name,
        'spec_col': spec_col,
        'curr_color': curr_color,
        'curr_blouse': curr_blouse,
        'top_colors': top_colors
    })

print(f"{'#':2s} | {'Row':3s} | {'SKU':20s} | {'Current Col':12s} | {'Spec Color':35s} | {'Top Colors in Image'}")
print("-" * 120)
for i, it in enumerate(items, 1):
    tc_str = ", ".join([f"RGB{c}:{cnt}" for c, cnt in it['top_colors'][:3]])
    print(f"{i:2d} | R{it['row']:2d} | {it['sku']:20s} | {it['curr_color']:12s} | {it['spec_col'][:35]:35s} | {tc_str}")
